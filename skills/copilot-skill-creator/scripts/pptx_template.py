"""Template: build a .pptx with asserting headlines and detail in the notes.

Copy into a generated skill and adapt. Use this only when the harness's own
file creation cannot give the control you need - a branded template, an exact
slide order, a layout that must not vary.

Obeys the rules in references/output-pptx.md: headlines assert rather than
label, six bullets maximum, supporting detail goes in the speaker notes.
"""

from pptx import Presentation
from pptx.util import Pt

# --- adapt from here -------------------------------------------------------
DECK_TITLE = "Q3 credit review"
SUBTITLE = "Prepared 2026-09-20 for the credit committee"

# Headline states the point, not the topic. "Margin fell 4pts on freight",
# never "Margin". Notes carry the method, caveats and secondary numbers.
SLIDES = [
    {
        "headline": "Q3 revenue grew 8%, all of it from EMEA",
        "bullets": ["EMEA +14%", "AMER flat", "APAC -2%"],
        "notes": "Growth is volume, not price. FX contributed 1.1pts.",
    },
    {
        "headline": "One borrower drives 38% of portfolio concentration",
        "bullets": ["Acme Ltd GBP 750k", "Next largest GBP 210k"],
        "notes": "Concentration limit is 35%. Breach discussed under covenants.",
    },
]
OUTPUT = "credit-review.pptx"
MAX_BULLETS = 6
# --- to here ---------------------------------------------------------------

prs = Presentation()

title_slide = prs.slides.add_slide(prs.slide_layouts[0])
title_slide.shapes.title.text = DECK_TITLE
title_slide.placeholders[1].text = SUBTITLE

for spec in SLIDES:
    bullets = spec["bullets"]
    if len(bullets) > MAX_BULLETS:
        raise ValueError(
            f"{spec['headline']!r} has {len(bullets)} bullets; "
            f"max is {MAX_BULLETS}. Split the slide rather than shrinking the text."
        )

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = spec["headline"]

    body = slide.placeholders[1].text_frame
    body.text = bullets[0]
    for line in bullets[1:]:
        para = body.add_paragraph()
        para.text = line
        para.level = 0
        para.font.size = Pt(18)

    slide.notes_slide.notes_text_frame.text = spec["notes"]

prs.save(OUTPUT)
print(f"wrote {OUTPUT}: {len(SLIDES) + 1} slides including title")
