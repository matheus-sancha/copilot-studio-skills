"""Template: build a .docx with fixed sections and real heading styles.

Copy into a generated skill and adapt. Use this only when the harness's own
file creation cannot give the control you need - a house template, exact
styles, a section order that must never vary.

Obeys the rules in references/output-docx.md: fixed sections always present,
real heading styles rather than bold text, conclusion first, no placeholders
in a delivered file.
"""

from docx import Document
from docx.shared import Pt

# --- adapt from here -------------------------------------------------------
TITLE = "Credit memo: Acme Ltd"
SUMMARY = "Recommend approval at the requested limit, subject to the covenant below."

# Every section appears every time. A section with nothing to say says so.
SECTIONS = [
    ("Borrower", "Acme Ltd, incorporated 2014, manufacturing."),
    ("Financials", "Revenue GBP 4.2m, EBITDA GBP 610k, both FY2025 reported."),
    ("Risks", "Customer concentration: top client is 38% of revenue."),
    ("Covenants", "Net debt / EBITDA below 3.0x, tested quarterly."),
    ("Recommendation", "Approve at GBP 750k."),
]
OUTPUT = "credit-memo.docx"
# --- to here ---------------------------------------------------------------

doc = Document()

doc.add_heading(TITLE, level=1)

# Conclusion first. Readers of internal documents rarely reach the end.
lead = doc.add_paragraph()
run = lead.add_run(SUMMARY)
run.bold = True
run.font.size = Pt(11)

for heading, body in SECTIONS:
    doc.add_heading(heading, level=2)          # a real style, not bold text
    doc.add_paragraph(body if body else "Insufficient data. Source not supplied.")

doc.save(OUTPUT)

empty = [h for h, b in SECTIONS if not b]
print(f"wrote {OUTPUT}: {len(SECTIONS)} sections"
      + (f"; {len(empty)} marked insufficient: {', '.join(empty)}" if empty else ""))
